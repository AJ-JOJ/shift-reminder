"""读班表：CSV（date,shift 两列）或 XLSX，得到「日期 -> 班次」。"""

from __future__ import annotations

import csv
import datetime as dt
import io
from dataclasses import dataclass
from pathlib import Path

WEEKDAY_CN = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
DATE_FORMATS = ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d", "%Y年%m月%d日", "%m/%d", "%m-%d")


@dataclass
class DayEntry:
    date: dt.date
    shift: str  # 空字符串表示这天没有班次（休息）

    @property
    def weekday_cn(self) -> str:
        return WEEKDAY_CN[self.date.weekday()]

    @property
    def is_off(self) -> bool:
        return not self.shift.strip()

    @property
    def summary(self) -> str:
        return self.shift if self.shift else "休息"


def parse_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in DATE_FORMATS:
        try:
            parsed = dt.datetime.strptime(text, fmt)
        except ValueError:
            continue
        if "%Y" in fmt:
            return parsed.date()
        return _infer_year(parsed.month, parsed.day)
    return None


def _infer_year(month: int, day: int) -> dt.date | None:
    """日期没写年份时（如 10/15），挑离今天最近的那一年。"""
    anchor = dt.date.today()
    best: dt.date | None = None
    best_distance: int | None = None
    for year in (anchor.year - 1, anchor.year, anchor.year + 1):
        try:
            candidate = dt.date(year, month, day)
        except ValueError:
            continue
        distance = abs((candidate - anchor).days)
        if best_distance is None or distance < best_distance:
            best, best_distance = candidate, distance
    return best


def _read_csv(path: Path, cfg: dict) -> list[list[str]]:
    raw = path.read_bytes()
    encodings = [cfg.get("csv_encoding")] if cfg.get("csv_encoding") else ["utf-8-sig", "gb18030", "utf-8"]
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
        except (UnicodeDecodeError, LookupError) as exc:
            last_error = exc
            continue
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
        rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if any(c.strip() for c in row)]
        if rows:
            return rows
    raise SystemExit(f"读不了这个 CSV（编码试过 {encodings}）：{last_error}")


def _read_xlsx(path: Path, cfg: dict) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("读 .xlsx 需要 openpyxl：pip install openpyxl")

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        wanted = str(cfg.get("sheet") or "").strip()
        titles = [wanted] if wanted else list(wb.sheetnames)
        if wanted and wanted not in wb.sheetnames:
            raise SystemExit(f"找不到工作表「{wanted}」")
        best_rows: list[list[str]] = []
        for title in titles:
            rows = [[("" if c is None else str(c).strip()) for c in row] for row in wb[title].iter_rows(values_only=True)]
            rows = [row for row in rows if any(row)]
            if len(rows) > len(best_rows):
                best_rows = rows
        return best_rows
    finally:
        wb.close()


def read_entries(schedule_path: Path, cfg: dict) -> list[DayEntry]:
    """返回按日期排好序的列表；每天一条（没有班次的那天 shift 为空字符串）。"""
    if not schedule_path.is_file():
        raise SystemExit(f"找不到班表文件：{schedule_path}")

    if schedule_path.suffix.lower() in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(schedule_path, cfg)
    else:
        rows = _read_csv(schedule_path, cfg)
    if not rows:
        raise SystemExit(f"班表是空的：{schedule_path}")

    # 认表头：第一行里有没有「日期」「班次」这类字样
    header = [c.strip().lower() for c in rows[0]]
    date_col, shift_col, header_found = 0, 1, False
    for idx, name in enumerate(header):
        if name in {"date", "日期", "时间", "班次日期"}:
            date_col, header_found = idx, True
        if name in {"shift", "班次", "班别", "排班", "值班"}:
            shift_col, header_found = idx, True
    data_rows = rows[1:] if header_found else rows

    entries: dict[dt.date, DayEntry] = {}
    skipped: list[str] = []
    for row in data_rows:
        need = max(date_col, shift_col) + 1
        if len(row) < need:
            row = list(row) + [""] * (need - len(row))
        day = parse_date(row[date_col])
        if day is None:
            if any(c.strip() for c in row):
                skipped.append(" | ".join(c for c in row if c.strip()))
            continue
        entries[day] = DayEntry(date=day, shift=(row[shift_col] or "").strip())

    if not entries:
        raise SystemExit("没从班表里读出任何日期。请检查：第一列是不是日期（如 2026/10/1），第二列是不是班次。")
    result = [entries[d] for d in sorted(entries)]
    if skipped:
        print(f"[提示] 有 {len(skipped)} 行没认出日期，已跳过（例如：{skipped[0]}）")
    return result


def find_entry(entries: list[DayEntry], target: dt.date) -> DayEntry | None:
    for entry in entries:
        if entry.date == target:
            return entry
    return None


def needs_reminder(entry: DayEntry, cfg: dict) -> tuple[bool, str]:
    """判断这天要不要提醒。

    - 空白 = 休息，不提醒；
    - skip_shifts 里的班次不提醒（例如「下大夜」是大夜班的延续日）；
    - 其余有班次的日子都提醒。
    用排除法而不是白名单，是因为「下大夜」里含有「大夜」，白名单会误判。
    """
    if entry is None:
        return False, "班表里没有这一天"
    if entry.is_off:
        return False, "这天没有班次（休息）"
    skip = [str(s).strip() for s in cfg.get("skip_shifts", []) if str(s).strip()]
    for keyword in skip:
        if keyword in entry.shift:
            return False, f"「{entry.shift}」是不提醒的班次"
    return True, ""
